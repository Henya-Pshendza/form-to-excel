"""
app/exporter.py
---------------
מייצא את תוצאות החילוץ לקובץ אקסל.

פורמט הפלט:
- שורה אחת לכל טופס התחייבות
- עמודה לכל שדה (שדות multi מקבלים עמודות נפרדות: מרכזי_1, מרכזי_2...)
- שדות "לא בטוח" מוצגים ברקע כתום
- שדות "חסר" מוצגים ברקע אדום
- שורת כותרת קבועה ומעוצבת
"""

from __future__ import annotations

from datetime import datetime
from pathlib import Path
from typing import Any

import openpyxl
from openpyxl.styles import (
    PatternFill, Font, Alignment, Border, Side
)
from openpyxl.utils import get_column_letter

from app.confidence_handler import FieldStatus, STATUS_COLORS
from config.fields import FIELDS
from utils.file_utils import ensure_output_dir
from utils.logger import get_logger

logger = get_logger(__name__)

# צבעי תאים — מחרוזות Hex ללא #
FILL_OK      = "D4EDDA"   # ירוק בהיר
FILL_LOW     = "FFF3CD"   # כתום בהיר
FILL_MISSING = "F8D7DA"   # אדום בהיר
FILL_ERROR   = "E2E3E5"   # אפור
FILL_HEADER  = "2C3E50"   # כהה לשורת כותרת

# רוחב עמודה ברירת מחדל
DEFAULT_COL_WIDTH = 20


def export_to_excel(
    forms_data: list[dict[str, Any]],
    filename: str | None = None,
) -> Path:
    """
    מייצא רשימת טפסים מעובדים לקובץ אקסל.

    Args:
        forms_data: רשימה של תוצאות pipeline — כל פריט הוא טופס אחד.
                    כל פריט: { שדה: {value, confidence, status}, "source_file": str }
        filename:   שם קובץ הפלט (ללא נתיב). אם None — נוצר שם עם timestamp.

    Returns:
        נתיב לקובץ האקסל שנוצר.
    """
    if not forms_data:
        raise ValueError("אין נתונים לייצוא")

    # בניית שם קובץ
    output_dir = ensure_output_dir()
    if filename is None:
        timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
        filename = f"התחייבויות_{timestamp}.xlsx"
    output_path = output_dir / filename

    logger.info(f"מייצא {len(forms_data)} טפסים ל-{output_path.name}")

    # בניית רשימת עמודות דינמית
    columns = _build_column_list(forms_data)

    # יצירת workbook
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "התחייבויות"
    ws.sheet_view.rightToLeft = True   # RTL לעברית

    # כתיבת שורת כותרת
    _write_header(ws, columns)

    # כתיבת שורת נתונים לכל טופס
    for row_idx, form in enumerate(forms_data, start=2):
        _write_form_row(ws, row_idx, form, columns)

    # עיצוב: קפיאת שורת כותרת, מסנן
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions

    # כיוון גיליון RTL
    ws.sheet_view.rightToLeft = True

    wb.save(str(output_path))
    logger.info(f"אקסל נשמר: {output_path}")
    return output_path


def _build_column_list(forms_data: list[dict]) -> list[str]:
    """
    בונה רשימת עמודות דינמית:
    - עמודת "קובץ מקור" ראשונה
    - לכל שדה רגיל — עמודה אחת
    - לשדות multi (מרכזי, התחייבות) — עמודות לפי המקסימום שנמצא בנתונים

    Args:
        forms_data: כל הטפסים

    Returns:
        רשימת שמות עמודות לפי הסדר
    """
    columns = ["קובץ מקור"]

    # סופר מקסימום ערכים בשדות multi
    multi_max: dict[str, int] = {}
    for field_key, field_def in FIELDS.items():
        if field_def["multi"]:
            max_count = 1
            for form in forms_data:
                val = form.get(field_key, {})
                if not isinstance(val, dict):
                    val = {}
                val = val.get("value", [])
                if isinstance(val, list):
                    max_count = max(max_count, len(val))
            multi_max[field_key] = max_count

    # בנייה לפי סדר FIELDS
    for field_key, field_def in FIELDS.items():
        if field_def["multi"]:
            count = multi_max.get(field_key, 1)
            base = field_def["excel_col"]
            if count == 1:
                columns.append(base)
            else:
                for i in range(1, count + 1):
                    columns.append(f"{base}_{i}")
        else:
            columns.append(field_def["excel_col"])

    return columns


def _write_header(ws: openpyxl.worksheet.worksheet.Worksheet, columns: list[str]) -> None:
    """
    כותב שורת כותרת מעוצבת.

    Args:
        ws:      גיליון האקסל
        columns: רשימת שמות העמודות
    """
    header_fill = PatternFill("solid", fgColor=FILL_HEADER)
    header_font = Font(bold=True, color="FFFFFF", size=11)
    header_align = Alignment(horizontal="center", vertical="center", wrap_text=True)
    thin_border = Border(
        left=Side(style="thin", color="FFFFFF"),
        right=Side(style="thin", color="FFFFFF"),
    )

    for col_idx, col_name in enumerate(columns, start=1):
        cell = ws.cell(row=1, column=col_idx, value=col_name)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = header_align
        cell.border = thin_border

        # רוחב עמודה
        col_letter = get_column_letter(col_idx)
        ws.column_dimensions[col_letter].width = DEFAULT_COL_WIDTH

    ws.row_dimensions[1].height = 30


def _write_form_row(
    ws:      openpyxl.worksheet.worksheet.Worksheet,
    row_idx: int,
    form:    dict[str, Any],
    columns: list[str],
) -> None:
    """
    כותב שורת נתונים אחת לטופס אחד.

    Args:
        ws:      גיליון האקסל
        row_idx: מספר השורה (מתחיל מ-2)
        form:    נתוני הטופס
        columns: רשימת שמות העמודות
    """
    # בניית מיפוי שם_עמודה → (ערך, סטטוס)
    cell_data: dict[str, tuple[Any, FieldStatus]] = {}

    # קובץ מקור
    cell_data["קובץ מקור"] = (form.get("source_file", ""), FieldStatus.OK)

    for field_key, field_def in FIELDS.items():
        field_info = form.get(field_key, {})
        # מגן מפני מקרה שהערך אינו dict
        if not isinstance(field_info, dict):
            field_info = {}
        value      = field_info.get("value")
        status     = field_info.get("status", FieldStatus.MISSING)

        if field_def["multi"]:
            base = field_def["excel_col"]
            items = value if isinstance(value, list) else []

            # מוצאים כמה עמודות יש לשדה הזה
            multi_cols = [c for c in columns if c == base or c.startswith(f"{base}_")]

            for i, col_name in enumerate(multi_cols):
                item_val = items[i] if i < len(items) else None
                # סטטוס: אם הרשימה ריקה — MISSING, אחרת הסטטוס הכללי
                item_status = status if items else FieldStatus.MISSING
                cell_data[col_name] = (item_val, item_status)
        else:
            cell_data[field_def["excel_col"]] = (value, status)

    # כתיבה לגיליון
    default_align = Alignment(horizontal="right", vertical="center", wrap_text=False)

    for col_idx, col_name in enumerate(columns, start=1):
        value, status = cell_data.get(col_name, (None, FieldStatus.MISSING))

        cell = ws.cell(row=row_idx, column=col_idx, value=_format_cell_value(value))
        cell.alignment = default_align
        cell.fill = _get_fill(status)

    ws.row_dimensions[row_idx].height = 18


def _format_cell_value(value: Any) -> Any:
    """
    ממיר ערך לפורמט מתאים לתא אקסל.
    רשימות → מחרוזת מופרדת בפסיקים.
    None → ריק.

    Args:
        value: ערך כלשהו

    Returns:
        ערך מוכן לתא אקסל
    """
    if value is None:
        return ""
    if isinstance(value, list):
        return ", ".join(str(v) for v in value if v)
    return value


def _get_fill(status: FieldStatus) -> PatternFill:
    """
    מחזיר PatternFill לפי סטטוס השדה.

    Args:
        status: FieldStatus

    Returns:
        PatternFill מתאים
    """
    color_map = {
        FieldStatus.OK:      FILL_OK,
        FieldStatus.LOW:     FILL_LOW,
        FieldStatus.MISSING: FILL_MISSING,
        FieldStatus.ERROR:   FILL_ERROR,
    }
    color = color_map.get(status, FILL_ERROR)
    return PatternFill("solid", fgColor=color)


def forms_to_export_list(flagged_forms: list[dict]) -> list[dict]:
    """
    ממיר רשימת טפסים עם confidence_flags לפורמט ייצוא שטוח.
    משמש כ-adapter בין pipeline לבין export_to_excel.

    Args:
        flagged_forms: רשימת תוצאות עם status לכל שדה

    Returns:
        אותה רשימה (כבר במבנה הנכון)
    """
    return flagged_forms
